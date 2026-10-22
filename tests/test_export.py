import json
from dataclasses import replace

import pytest
from docx import Document

from qpgen.config import load_config
from qpgen.export import export, paper_json
from qpgen.paper import build_paper, spec_from_pattern, spec_from_preset
from tests.test_paper import TOPICS, stub_generate

CFG = load_config()


@pytest.fixture
def paper():
    return build_paper(spec_from_preset(CFG, "DBMS", TOPICS, "balanced"), 42, CFG, stub_generate)


@pytest.fixture
def out_cfg(tmp_path):
    return replace(CFG, paths={**CFG.paths, "outputs_dir": str(tmp_path)})


def test_writes_three_files_with_expected_names(paper, out_cfg, tmp_path):
    paths = export(paper, out_cfg, "dbms", ["docx", "pdf", "json"], stamp="20261012_1030")
    assert [p.name for p in paths] == ["20261012_1030_dbms.docx", "20261012_1030_dbms.pdf", "20261012_1030_dbms.json"]
    assert all(p.parent == tmp_path and p.stat().st_size > 0 for p in paths)


def test_docx_opens_and_has_layout(paper, out_cfg):
    path = export(paper, out_cfg, "dbms", ["docx"], stamp="t")[0]
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "MMCOE" in text and "Section A (10 x 2 = 20 marks)" in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert cells.count("[2]") == 10 and cells.count("[5]") == 6 and cells.count("[10]") == 2
    assert cells[:4] == ["Question", "CO", "BL", "Marks"]
    assert cells[4].startswith("1. ") and cells[5] == "CO1" and cells[6].startswith("BL")
    assert not any("generated" in c for c in cells)
    assert doc.styles["Normal"].font.name == "Times New Roman"


def test_any_n_of_m_title(out_cfg):
    spec = spec_from_pattern(CFG, "DBMS", TOPICS, [{"marks": 5, "count": 6, "attempt": 4}])
    doc = Document(export(build_paper(spec, 1, CFG, stub_generate), out_cfg, "dbms", ["docx"], stamp="t")[0])
    assert "Section A (Answer any 4 of 6, 4 x 5 = 20 marks)" in "\n".join(p.text for p in doc.paragraphs)


def test_pdf_is_a_pdf(paper, out_cfg):
    path = export(paper, out_cfg, "dbms", ["pdf"], stamp="t")[0]
    data = path.read_bytes()
    assert data[:5] == b"%PDF-" and b"Times-Roman" in data and b"Times-BoldItalic" in data


def test_json_has_spec_seed_and_scores(paper, out_cfg):
    data = json.loads(export(paper, out_cfg, "dbms", ["json"], stamp="t")[0].read_text(encoding="utf-8"))
    assert data["seed"] == 42 and data["spec"]["institute"] == "MMCOE"
    q = data["sections"][0]["questions"][0]
    assert {"origin", "relevance", "difficulty_prob", "text", "marks"} <= set(q)
    assert sum(q["marks"] for s in data["sections"] for q in s["questions"]) == data["spec"]["total_marks"]
    assert {"co", "bloom"} <= set(q) and sum(data["co_marks"].values()) == 70


def test_same_paper_same_json(paper):
    again = build_paper(spec_from_preset(CFG, "DBMS", TOPICS, "balanced"), 42, CFG, stub_generate)
    assert paper_json(paper) == paper_json(again)
