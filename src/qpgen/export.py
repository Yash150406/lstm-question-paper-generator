"""Write a paper as DOCX, PDF and a JSON sidecar."""
import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from qpgen.config import Config
from qpgen.paper import co_marks

HEADERS = ["Question", "CO", "BL", "Marks"]
WIDTHS = [12.8, 1.4, 1.3, 1.5]   # cm, fits A4 with 2 cm margins


def numbered(paper: dict) -> list[tuple[dict, list[tuple[int, dict]]]]:
    """Sections with questions numbered straight through the paper."""
    n, out = 0, []
    for s in paper["sections"]:
        rows = []
        for q in s["questions"]:
            n += 1
            rows.append((n, q))
        out.append((s, rows))
    return out


def header_lines(paper: dict) -> tuple[str, str, str]:
    """Institute, exam title with subject, and the duration and marks line."""
    spec = paper["spec"]
    return (spec.institute, f"{spec.exam_title}: {spec.subject}",
            f"Duration: {spec.duration_minutes} minutes  |  Maximum marks: {spec.total_marks}")


def section_title(s: dict) -> str:
    """Section A (10 x 2 = 20 marks), or Section B (Answer any 4 of 6, 4 x 5 = 20 marks)."""
    n, k = len(s["questions"]), s.get("attempt") or len(s["questions"])
    choice = f"Answer any {k} of {n}, " if k < n else ""
    return f"Section {s['name']} ({choice}{k} x {s['marks']} = {k * s['marks']} marks)"


def row(n: int, q: dict) -> list[str]:
    """Question, CO, BL and marks cells for one question."""
    return [f"{n}. {q['text']}", q.get("co", ""), q.get("bloom", ""), f"[{q['marks']}]"]


def write_docx(paper: dict, path: Path) -> Path:
    """Header block, instructions, section headings and numbered questions with marks on the right."""
    doc = Document()
    for i, line in enumerate(header_lines(paper)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(line)
        run.bold = i < 2
        run.font.size = Pt(14 if i == 0 else 11)
    for s, rows in numbered(paper):
        heading = doc.add_heading(section_title(s), level=2)
        # Keep a section heading on the same page as its first question.
        heading.paragraph_format.keep_with_next = True
        if s["instruction"] not in section_title(s):
            note = doc.add_paragraph(s["instruction"])
            note.runs[0].italic = note.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=0, cols=len(HEADERS))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, values in enumerate([HEADERS] + [row(n, q) for n, q in rows]):
            cells = table.add_row().cells
            for j, (cell, value, width) in enumerate(zip(cells, values, WIDTHS)):
                cell.width, cell.text = Cm(width), value
                para = cell.paragraphs[0]
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.keep_with_next = i < 2
                if i == 0:
                    para.runs[0].bold = True
    doc.save(path)
    return path


def write_pdf(paper: dict, path: Path) -> Path:
    """Same layout as the DOCX, built with reportlab."""
    styles = getSampleStyleSheet()
    center = ParagraphStyle("center", parent=styles["Normal"], alignment=1, fontSize=11, leading=15)
    title = ParagraphStyle("title", parent=center, fontName="Helvetica-Bold", fontSize=14, leading=18)
    inst, exam, info = header_lines(paper)
    story = [Paragraph(inst, title), Paragraph(f"<b>{exam}</b>", center), Paragraph(info, center), Spacer(1, 12)]
    for s, rows in numbered(paper):
        story.append(Paragraph(section_title(s), styles["Heading3"]))
        if s["instruction"] not in section_title(s):
            story.append(Paragraph(f"<i>{s['instruction']}</i>", styles["Normal"]))
        data = [HEADERS] + [[Paragraph(r[0], styles["Normal"]), *r[1:]] for r in (row(n, q) for n, q in rows)]
        table = Table(data, colWidths=[w * cm for w in WIDTHS], repeatRows=1)
        table.setStyle(TableStyle([("ALIGN", (1, 0), (-1, -1), "CENTER"), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                                   ("LINEBELOW", (0, 0), (-1, 0), 0.5, colors.black),
                                   ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
        story += [table, Spacer(1, 8)]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                      topMargin=2 * cm, bottomMargin=2 * cm, title=exam).build(story)
    return path


def paper_json(paper: dict) -> str:
    """Spec, seed, marks per course outcome and every question with its origin and scores.
    No timestamps, so reruns match."""
    return json.dumps({"spec": asdict(paper["spec"]), "seed": paper["seed"], "co_marks": co_marks(paper),
                       "sections": paper["sections"]}, indent=2, ensure_ascii=False)


def write_json(paper: dict, path: Path) -> Path:
    """Write the JSON sidecar."""
    path.write_text(paper_json(paper), encoding="utf-8")
    return path


WRITERS = {"docx": write_docx, "pdf": write_pdf, "json": write_json}


def export(paper: dict, cfg: Config, subject_key: str, formats: list[str], stamp: str | None = None) -> list[Path]:
    """Write the requested formats to outputs/<YYYYMMDD_HHMM>_<subject>.<ext>."""
    out = cfg.path("outputs_dir")
    out.mkdir(parents=True, exist_ok=True)
    stamp = stamp or datetime.now().strftime("%Y%m%d_%H%M")
    return [WRITERS[f](paper, out / f"{stamp}_{subject_key}.{f}") for f in formats]
